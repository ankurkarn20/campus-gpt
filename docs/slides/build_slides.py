"""Build the CampusGPT major-project presentation on the department's slide template.

Usage (from the repository root):
    python docs/slides/render_diagrams.py      # diagrams -> docs/slides/build/*.png
    python docs/slides/build_slides.py TEMPLATE.pptx

TEMPLATE.pptx is the sample deck supplied by the department (UPES layouts, theme and
logos). Its slides are discarded; only its masters, layouts, theme and cover images
are reused. Output: docs/slides/CampusGPT_Major_Project.pptx
"""
import re
import sys
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "build"
OUT = HERE / "CampusGPT_Major_Project.pptx"

TEXT = MSO_THEME_COLOR.TEXT_1
WHITE = MSO_THEME_COLOR.BACKGROUND_1
BLUE = MSO_THEME_COLOR.ACCENT_2          # template heading blue (#1F5E99)
STEP_COLOURS = [MSO_THEME_COLOR.ACCENT_2, MSO_THEME_COLOR.ACCENT_3, MSO_THEME_COLOR.ACCENT_4,
                MSO_THEME_COLOR.ACCENT_5, MSO_THEME_COLOR.ACCENT_6, MSO_THEME_COLOR.TEXT_2]

LEFT = 0.42          # content left edge (inches), as in the template
WIDTH = 12.5         # content width (inches)


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def _runs(p, text, size, bold=False, color=TEXT, font=None, italic=False):
    """Add runs to paragraph p; **double asterisks** mark bold segments."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italic
        r.font.color.theme_color = color
        if font:
            r.font.name = font


def _bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def textbox(slide, x, y, w, h, paras, size=18, bullets=False, bold=False, color=TEXT,
            align=None, space_after=None, font=None, anchor=None, italic=False):
    """paras: list of strings (or a single string). Returns the text frame."""
    if isinstance(paras, str):
        paras = [paras]
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if space_after is not None:
            p.space_after = Pt(space_after)
        if bullets:
            _bullet(p)
        if align:
            p.alignment = align
        _runs(p, text, size, bold=bold, color=color, font=font, italic=italic)
    return tf


def heading(slide, x, y, w, text, size=22):
    textbox(slide, x, y, w, 0.45, text, size=size, bold=True, color=BLUE)


def rect(slide, x, y, w, h, fill=WHITE, line=BLUE, line_pt=1.25, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.theme_color = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.theme_color = line
        s.line.width = Pt(line_pt)
    s.text_frame.text = ""
    return s


def picture(slide, name, x, y, max_w, max_h, align="center"):
    """Place docs/slides/build/<name>.png inside the box, keeping its aspect ratio."""
    path = IMG / f"{name}.png"
    with Image.open(path) as im:
        ratio = im.width / im.height
    w, h = max_w, max_w / ratio
    if h > max_h:
        h, w = max_h, max_h * ratio
    left = x + (max_w - w) / 2 if align == "center" else x
    slide.shapes.add_picture(str(path), Inches(left), Inches(y), Inches(w), Inches(h))
    return left, w, h


# ---------------------------------------------------------------------------
# Table helper (plain black grid, light-blue header -- as in the template)
# ---------------------------------------------------------------------------
def _cell_style(cell, header):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(edge), w="12700", cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:schemeClr"), val="tx1")
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    fill = etree.SubElement(tcPr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr"), val="bg2" if header else "bg1")
    tcPr.set("marL", "91440")
    tcPr.set("marR", "91440")
    tcPr.set("marT", "36576")
    tcPr.set("marB", "36576")
    tcPr.set("anchor", "ctr")


def table(slide, x, y, col_widths, rows, size=12.5, row_h=0.3):
    n_rows, n_cols = len(rows), len(col_widths)
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(sum(col_widths)), Inches(row_h * n_rows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            _cell_style(cell, header=(i == 0))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _runs(p, value, size, bold=(i == 0))
    return gf


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------
def content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("MP_CONTENT"))
    ph = next(iter(slide.placeholders))
    ph.text_frame.text = ""
    _runs(ph.text_frame.paragraphs[0], title, 32, bold=True)
    return slide


def cover(prs, logos):
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("MP_COVER"))
    for blob, (x, y, w, h) in logos:
        slide.shapes.add_picture(BytesIO(blob), x, y, w, h)
    textbox(slide, 0.40, 2.55, 12.5, 0.9, "Major Project", size=54, align=PP_ALIGN.CENTER)
    textbox(slide, 0.60, 3.60, 12.1, 1.3,
            "Title: CampusGPT: AI-Powered University Intelligence and Student "
            "Assistance Platform", size=30, align=PP_ALIGN.CENTER)
    tf = textbox(slide, 0.40, 5.15, 6.5, 2.0, "Presented by:", size=18, bold=True)
    for name in ["Ankur Karn  500118986", "Ashmita Sharma  500122422",
                 "Ananya Singh  500123837", "Ayushi Mishra  500121986"]:
        _runs(tf.add_paragraph(), name, 18)
    tf = textbox(slide, 8.20, 5.15, 4.9, 2.0, "Institution:", size=18, bold=True)
    for line in ["School of Computer Science, UPES Dehradun", "B.Tech CSE (Cyber Security), Batch 11",
                 "Mentor: Mr. Yogesh Rathia", "October 2026"]:
        _runs(tf.add_paragraph(), line, 18)


def steps(slide, items, y=1.3):
    """Row of six coloured step boxes, as on the template's methodology slide."""
    for i, (title, desc) in enumerate(items):
        x = LEFT + i * 2.10
        rect(slide, x, y, 1.95, 1.9, fill=STEP_COLOURS[i], line=None)
        circle = rect(slide, x + 0.70, y + 0.10, 0.55, 0.55, fill=WHITE, line=WHITE, line_pt=1,
                      shape=MSO_SHAPE.OVAL)
        tf = circle.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _runs(p, str(i + 1), 18, bold=True, color=STEP_COLOURS[i])
        textbox(slide, x + 0.05, y + 0.70, 1.85, 0.4, title, size=17, bold=True, color=WHITE,
                align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.10, y + 1.15, 1.75, 0.7, desc, size=14, color=WHITE,
                align=PP_ALIGN.CENTER)


def panel(slide, y, h, title, body, size=20):
    rect(slide, LEFT, y, WIDTH, h)
    textbox(slide, LEFT + 0.23, y + 0.10, WIDTH - 0.5, 0.45, title, size=22, bold=True,
            color=BLUE)
    textbox(slide, LEFT + 0.23, y + 0.60, WIDTH - 0.5, h - 0.7, body, size=size)


# ---------------------------------------------------------------------------
# Deck
# ---------------------------------------------------------------------------
SECTIONS = [
    "Abstract", "Introduction & Problem Statement", "Objectives", "Literature Review",
    "Dataset Description", "Methodology", "System Architecture", "Data Flow Diagram",
    "Inter-Process Communication", "Database Design", "Implementation: Concepts and Libraries",
    "Security and Data Isolation", "SWOT Analysis", "Project Schedule",
    "Limitations, Future Work & Conclusion", "References",
]


def build(template: Path):
    prs = Presentation(str(template))

    # Keep the cover logos, then drop every sample slide.
    logos = [(sh.image.blob, (sh.left, sh.top, sh.width, sh.height))
             for sh in prs.slides[0].shapes if sh.shape_type == 13]
    for sld_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld_id.rId)
        prs.slides._sldIdLst.remove(sld_id)

    n = iter(range(1, 100))
    title = lambda name: f"{next(n)}. {name}"

    # Cover ----------------------------------------------------------------
    cover(prs, logos)

    # Contents -------------------------------------------------------------
    s = content_slide(prs, "Content")
    half = (len(SECTIONS) + 1) // 2
    for col, chunk in enumerate([SECTIONS[:half], SECTIONS[half:]]):
        start = 1 + col * half
        textbox(s, 0.75 + col * 5.6, 1.15, 5.4, 5.9,
                [f"{start + i}. {name}" for i, name in enumerate(chunk)], size=17,
                space_after=8)

    # 1. Abstract ------------------------------------------------------------
    s = content_slide(prs, title("Abstract"))
    textbox(s, LEFT, 1.25, WIDTH, 5.8, [
        "Students search **several portals, PDFs and notice boards** for attendance, exams, "
        "fees and rules, and general chatbots **cannot see their records**.",
        "We build **CampusGPT**, an AI helpdesk that answers in plain language from the "
        "student's **own records** and the university's **policy documents**.",
        "The AI never touches the database: it calls **9 read-only tools**, and the backend "
        "decides **whose data is returned** from the login session.",
        "Policy answers use **hybrid retrieval (RAG)** with **verified citations**; "
        "calculations such as attendance shortfall are done by **deterministic code**.",
        "Built on a **5,000-student synthetic database** with a free stack (FastAPI, "
        "PostgreSQL + pgvector, Next.js, free-tier LLM) and evaluated against "
        "**three simpler baselines**.",
    ], size=22, bullets=True, space_after=14)

    # 2. Introduction & Problem Statement -------------------------------------
    s = content_slide(prs, title("Introduction & Problem Statement"))
    for i, (num, label) in enumerate([("5,000", "synthetic student records"),
                                      ("9", "permission-checked AI tools"),
                                      ("0", "direct database access for the AI")]):
        x = LEFT + i * 4.2
        textbox(s, x, 1.10, 4.0, 0.7, num, size=36, bold=True, color=BLUE,
                align=PP_ALIGN.CENTER)
        textbox(s, x, 1.80, 4.0, 0.35, label, size=14, align=PP_ALIGN.CENTER)
    textbox(s, LEFT, 2.40, WIDTH, 2.3, [
        "**Scattered information**: attendance in the ERP, rules in PDFs, deadlines on "
        "notice boards.",
        "**Generic chatbots** know nothing about the student and can invent policy "
        "(hallucination).",
        "**Privacy risk**: giving an AI direct database access could expose another "
        "student's data.",
    ], size=20, bullets=True, space_after=10)
    panel(s, 4.85, 2.2, "Problem Statement",
          "Students have no single place to ask questions that need both **their own "
          "records** and **university rules**, and existing AI assistants are either "
          "**ungrounded** or **unsafe** with personal data. There is a need for a **free, "
          "grounded and secure assistant** that answers from verified sources and never "
          "reveals one student's data to another.")

    # 3. Objectives ----------------------------------------------------------
    s = content_slide(prs, title("Objectives"))
    heading(s, LEFT, 1.10, WIDTH, "Main Objective")
    textbox(s, LEFT, 1.55, WIDTH, 0.9,
            "Design and build a secure AI assistant that answers students' academic and "
            "administrative questions from their own records and university policy "
            "documents, with sources, without giving the AI direct database access.",
            size=18)
    heading(s, LEFT, 2.60, WIDTH, "Sub-Objectives")
    textbox(s, LEFT, 3.05, WIDTH, 4.0, [
        "**Data**: a reproducible synthetic university of 5,000 students with consistent "
        "courses, attendance, exams and fees.",
        "**Security**: login, student and admin roles, and backend-enforced data isolation.",
        "**AI tools**: 9 read-only tools with schema validation and audit logging.",
        "**Retrieval**: hybrid vector + keyword search over policy documents with verified "
        "citations.",
        "**Campus intelligence**: a prioritised “what to take care of this week” "
        "summary.",
        "**Evaluation**: answer and citation accuracy, tool selection, latency and leakage "
        "tests against baselines.",
    ], size=20, bullets=True, space_after=8)

    # 4. Literature Review ---------------------------------------------------
    s = content_slide(prs, title("Literature Review"))
    table(s, LEFT, 1.15, [1.55, 3.05, 2.6, 2.75, 2.55], [
        ["Year - Ref.", "Title", "Proposed Work", "Results", "Limitations"],
        ["2009 - Cormack et al. [1]", "Reciprocal rank fusion outperforms Condorcet and "
         "individual rank learning methods", "Merge ranked lists by 1/(k + rank)",
         "Simple fusion beat more complex learned methods", "Fixed constant k; no learning"],
        ["2020 - Lewis et al. [2]", "Retrieval-augmented generation for knowledge-intensive "
         "NLP tasks", "Generator conditioned on passages from a dense retriever",
         "More factual answers than parametric-only models on open-domain QA",
         "General knowledge only; no access control"],
        ["2023 - Yao et al. [3]", "ReAct: Synergizing reasoning and acting in language models",
         "Interleave reasoning steps with tool actions",
         "Beat reasoning-only and acting-only prompting on QA and decision tasks",
         "Tools are unrestricted; errors compound"],
        ["2023 - Schick et al. [4]", "Toolformer: Language models can teach themselves to "
         "use tools", "Model learns when to call APIs such as a calculator or search",
         "Improved zero-shot results on several tasks", "No permission model for tools"],
        ["2023 - Greshake et al. [5]", "Compromising real-world LLM-integrated applications "
         "with indirect prompt injection", "Instructions hidden in retrieved content hijack "
         "the model", "Attacks shown on real LLM-integrated applications",
         "Prompts alone are not a safeguard"],
        ["2023 - Ji et al. [6]", "Survey of hallucination in natural language generation",
         "Taxonomy of hallucination and its mitigations",
         "Grounding in retrieved evidence is a key mitigation",
         "No single method removes hallucination"],
        ["2024 - Es et al. [7]", "RAGAS: Automated evaluation of retrieval augmented "
         "generation", "Reference-free metrics: faithfulness, answer and context relevance",
         "Evaluates RAG answers without labelled references", "Relies on an LLM as judge"],
    ], size=12, row_h=0.6)
    textbox(s, LEFT, 6.55, WIDTH, 0.45,
            "**Gap:** none of these combines personal records, policy retrieval and "
            "backend-enforced access control, which is the focus of CampusGPT.",
            size=15, color=BLUE)

    # 5. Dataset Description -------------------------------------------------
    s = content_slide(prs, title("Dataset Description"))
    textbox(s, LEFT, 1.25, 6.2, 5.7, [
        "**Synthetic by design**: no real university data; students, notices and "
        "documents are labelled is_synthetic.",
        "**Scale**: 5,000 students, 148 faculty, 190 courses and about 825,000 "
        "attendance records, generated and loaded in about 20 seconds.",
        "**Realistic**: Indian names (Faker en_IN); program, semester and courses always "
        "consistent; deliberate edge cases such as attendance just below 75%.",
        "**Reproducible**: fixed random seed; one command rebuilds the database.",
        "**Policy documents**: team-written rules for attendance, exams, fees and hostel, "
        "clearly marked fictional.",
    ], size=18, bullets=True, space_after=10)
    table(s, 6.85, 1.35, [1.75, 2.95, 1.4], [
        ["Table", "Contents", "Rows"],
        ["students", "Name, program, semester, section", "5,000"],
        ["enrollments", "Student ↔ course offering", "25,000"],
        ["class_sessions", "One row per class held", "15,338"],
        ["attendance", "Present / absent per session", "824,636"],
        ["assignments", "Title, due date, marks", "1,860"],
        ["submissions", "Submitted at, marks", "65,378"],
        ["exams", "Type, date, room", "1,395"],
        ["invoices", "Term, amount, due date, status", "10,000"],
    ], size=12.5, row_h=0.5)

    # 6. Methodology ---------------------------------------------------------
    s = content_slide(prs, title("Methodology"))
    steps(s, [("Data", "Synthetic DB, 5,000 students"),
              ("Secure", "Login, roles, data isolation"),
              ("Tools", "9 read-only tools"),
              ("Retrieve", "Hybrid search + citations"),
              ("Orchestrate", "LLM tool loop via adapter"),
              ("Serve", "Chat + dashboard (Next.js)")])
    heading(s, LEFT, 3.50, 6.0, "Answering a Question")
    textbox(s, LEFT, 4.05, 6.1, 2.9, [
        "The LLM chooses tools; the backend checks the role and validates arguments",
        "Tools run with the **session's** student ID",
        "Policy chunks are passed to the LLM as untrusted text",
        "Every citation is checked against the retrieved chunks",
    ], size=18, bullets=True, space_after=8)
    heading(s, 6.90, 3.50, 6.0, "Evaluation Metrics")
    textbox(s, 6.90, 4.05, 6.0, 2.9, [
        "Answer and citation accuracy on a 50-question benchmark",
        "Tool-selection accuracy and response latency",
        "Cross-student leakage and prompt-injection tests",
        "Comparison with three simpler baselines",
    ], size=18, bullets=True, space_after=8)

    # 7. System Architecture -------------------------------------------------
    s = content_slide(prs, title("System Architecture"))
    picture(s, "architecture", LEFT, 1.15, 8.4, 5.95, align="left")
    textbox(s, 9.05, 1.25, 3.95, 5.8, [
        "**Client**: Next.js student portal and admin console.",
        "**Application**: one FastAPI process with authentication, chat orchestrator, tool "
        "registry and retrieval.",
        "**Worker**: separate process for document ingestion.",
        "**Data**: PostgreSQL 16 + pgvector for records, vectors, job queue and logs.",
        "**(1)–(4)**: channels that cross a process boundary.",
    ], size=16, bullets=True, space_after=10)

    # 8. Data Flow Diagram ---------------------------------------------------
    s = content_slide(prs, title("Data Flow Diagram (Level 1)"))
    picture(s, "dfd-level1", LEFT, 1.15, 8.9, 5.95, align="left")
    textbox(s, 9.5, 1.25, 3.5, 5.8, [
        "**5 processes**, **5 data stores**, 3 external entities.",
        "Every flow passes through a process; no entity touches a data store directly.",
        "Student identity enters only through **1.0 Authenticate**.",
        "**3.0** reads only the logged-in student's rows.",
        "**5.0** runs in the separate worker process.",
    ], size=16, bullets=True, space_after=10)

    # 9. Inter-Process Communication -------------------------------------------
    s = content_slide(prs, title("Inter-Process Communication"))
    picture(s, "ipc-processes", LEFT, 1.15, 7.3, 4.45, align="left")
    picture(s, "job-states", LEFT + 0.6, 5.65, 5.0, 1.35, align="left")
    textbox(s, 7.95, 1.20, 5.0, 5.9, [
        "**(1)** REST + SSE streaming: browser ↔ API",
        "**(2)** PostgreSQL protocol, pooled: API / worker ↔ DB",
        "**(3)** LISTEN / NOTIFY: wakes the worker instantly",
        "**(4)** HTTPS: API → LLM provider",
        "**(5)** Shared job table with FOR UPDATE SKIP LOCKED: each job claimed once",
        "**(6)** HTTP: pages from the Next.js server",
        "**(7)** POSIX signals: graceful shutdown",
        "No extra message broker: PostgreSQL provides durability, locking and "
        "publish–subscribe.",
    ], size=15, bullets=True, space_after=6)

    # 10. Database Design --------------------------------------------------------
    s = content_slide(prs, title("Database Design"))
    picture(s, "er-academic", LEFT, 1.15, 7.4, 5.95, align="left")
    picture(s, "er-platform", 7.95, 1.20, 5.0, 3.1, align="left")
    textbox(s, 7.95, 4.45, 5.0, 2.6, [
        "**23 tables** in third normal form",
        "Derived values (attendance %) computed by a **view**, never stored",
        "Foreign keys, UNIQUE and CHECK constraints",
        "**HNSW** index for vectors, **GIN** for full-text",
    ], size=15, bullets=True, space_after=4)

    # 11. Implementation: Concepts and Libraries -------------------------------------
    s = content_slide(prs, title("Implementation: Concepts and Libraries"))
    heading(s, LEFT, 1.10, 6.0, "Programming Concepts")
    textbox(s, LEFT, 1.60, 6.2, 5.4, [
        "**OOP & polymorphism**: LLMProvider interface with Groq, Gemini and Ollama "
        "subclasses.",
        "**Design patterns**: Adapter, Strategy, Registry, Repository, Pub–Sub.",
        "**Async & generators**: async/await I/O; SSE answer streaming.",
        "**Decorators & DI**: @tool registry; FastAPI Depends(current_user).",
        "**Transactions & locking**: ACID writes; FOR UPDATE SKIP LOCKED.",
        "**Exact arithmetic**: Fraction for attendance calculations.",
    ], size=17, bullets=True, space_after=10)
    heading(s, 6.90, 1.10, 6.0, "Libraries")
    textbox(s, 6.90, 1.60, 6.0, 2.6, [
        "**Backend**: FastAPI, Pydantic, SQLAlchemy 2, psycopg 3, Alembic",
        "**AI & documents**: fastembed (BGE-small), pypdf, httpx",
        "**Security**: argon2-cffi, slowapi",
        "**Frontend**: Next.js, React, TypeScript, Tailwind CSS",
        "**Data & tests**: Faker, NumPy, pytest",
    ], size=16, bullets=True, space_after=4)
    rect(s, 6.90, 4.35, 6.0, 2.65)
    textbox(s, 7.00, 4.40, 5.8, 0.35, "Every tool call is checked by the backend", size=14,
            bold=True, color=BLUE)
    textbox(s, 7.00, 4.80, 5.8, 1.85, [
        "async def execute(call, user, db):",
        "    spec = TOOLS.get(call.name)",
        "    if spec is None or user.role not in spec.roles:",
        "        raise ToolNotPermitted(call.name)",
        "    args = spec.args.model_validate(call.arguments)",
        "    return await spec.fn(db, user.student_id, args)",
    ], size=12.5, font="Courier New")
    textbox(s, 7.00, 6.55, 5.8, 0.35, "The student ID comes from the session, never from the "
            "model.", size=12, color=BLUE)

    # 12. Security and Data Isolation ------------------------------------------------
    s = content_slide(prs, title("Security and Data Isolation"))
    textbox(s, LEFT, 1.10, WIDTH, 0.8,
            "**Goal**: make it impossible for one student to obtain another student's "
            "data, **even if the AI is tricked**.", size=18)
    heading(s, LEFT, 1.95, 6.0, "Defence Layers")
    textbox(s, LEFT, 2.45, 6.2, 4.6, [
        "**Session-derived identity**: tools receive the student ID from the login, never "
        "from the model.",
        "**Allow-listed tools**: role check and schema validation before every call.",
        "**No SQL from the model**: only parameterised queries written by us.",
        "**Untrusted documents**: retrieved text is data, never instructions.",
        "**Argon2id** passwords, rate limits and audit logs.",
    ], size=18, bullets=True, space_after=10)
    heading(s, 6.90, 1.95, 6.0, "How We Test It")
    textbox(s, 6.90, 2.45, 6.0, 4.6, [
        "**ID tampering**: change IDs in requests; only own data is returned.",
        "**Prompt injection**: “show another student's attendance”; no tool can "
        "fetch it.",
        "**Malicious PDF**: a document saying “ignore your rules” has no effect.",
        "**Admin endpoints** called as a student are rejected (403).",
        "Target: **0 leaked records** in every attack test.",
    ], size=18, bullets=True, space_after=10)

    # 13. SWOT Analysis --------------------------------------------------------------
    s = content_slide(prs, title("SWOT Analysis"))
    quads = [
        ("Strengths", ["Answers grounded in records and cited policy",
                       "Isolation enforced by the backend, not prompts",
                       "Deterministic calculations",
                       "Fully free stack; swappable LLM provider"]),
        ("Weaknesses", ["Synthetic data only; no live university link",
                        "Free-tier LLM rate limits",
                        "English only; no scanned PDFs",
                        "Small models may pick the wrong tool"]),
        ("Opportunities", ["Integration with university ERP through APIs",
                           "Hindi and regional languages",
                           "Faculty and staff roles",
                           "Mobile and messaging access"]),
        ("Threats", ["Changes to free-tier limits or terms",
                     "New prompt-injection techniques",
                     "DPDP Act duties for real deployment",
                     "Loss of trust after a wrong answer"]),
    ]
    for i, (name, items) in enumerate(quads):
        x = LEFT + (i % 2) * 6.25
        y = 1.15 + (i // 2) * 2.95
        rect(s, x, y, 6.25, 2.95, line=TEXT, line_pt=1)
        textbox(s, x + 0.15, y + 0.08, 5.9, 0.45, name, size=20, bold=True)
        textbox(s, x + 0.15, y + 0.58, 5.9, 2.3, items, size=18, bullets=True, space_after=4)

    # 14. Project Schedule -----------------------------------------------------------
    s = content_slide(prs, title("Project Schedule"))
    picture(s, "gantt", LEFT, 1.15, 8.6, 5.95, align="left")
    textbox(s, 9.3, 1.25, 3.7, 5.8, [
        "**31 days**: 13 Oct – 12 Nov 2026",
        "**Critical path**: A → B → D → E → G → H → J "
        "→ M → N",
        "Policy documents (C) and benchmark questions (L) run in parallel, with 9–12 "
        "days of slack",
        "Full PERT table (ES, EF, LS, slack) in the SRS",
    ], size=16, bullets=True, space_after=10)

    # 15. Limitations, Future Work & Conclusion -----------------------------------------
    s = content_slide(prs, title("Limitations, Future Work & Conclusion"))
    heading(s, LEFT, 1.10, 6.0, "Limitations")
    textbox(s, LEFT, 1.60, 6.2, 3.0, [
        "**Synthetic data only**: no authorised university access yet.",
        "**Free-tier LLM**: rate limits and provider changes.",
        "**English only**; scanned PDFs not supported.",
        "**Read-only tools**: students cannot submit requests.",
    ], size=19, bullets=True, space_after=8)
    heading(s, 6.90, 1.10, 6.0, "Future Work")
    textbox(s, 6.90, 1.60, 6.0, 3.0, [
        "**University API adapter** for real records.",
        "**Hindi and regional languages**.",
        "**Faculty role** and confirmation-gated actions (e.g. hostel leave).",
        "**Mobile and messaging** access.",
    ], size=19, bullets=True, space_after=8)
    panel(s, 4.75, 2.3, "Conclusion",
          "CampusGPT brings a student's own records and university policy into one "
          "conversational interface, with **answers grounded in verified sources** and "
          "**privacy enforced by the backend, not by the prompt**. The design is complete; "
          "implementation follows the 31-day plan, and evaluation will measure accuracy, "
          "citations and leakage against simpler baselines.", size=18)

    # 16. References ------------------------------------------------------------------
    s = content_slide(prs, title("References"))
    textbox(s, LEFT, 1.15, WIDTH, 5.9, [
        "[1] Cormack, G. V., Clarke, C. L. A., & Buettcher, S. (2009). Reciprocal rank fusion "
        "outperforms Condorcet and individual rank learning methods. Proceedings of the 32nd "
        "ACM SIGIR Conference, 758–759.",
        "[2] Lewis, P., et al. (2020). Retrieval-augmented generation for knowledge-intensive "
        "NLP tasks. Advances in Neural Information Processing Systems, 33, 9459–9474.",
        "[3] Yao, S., et al. (2023). ReAct: Synergizing reasoning and acting in language "
        "models. International Conference on Learning Representations (ICLR).",
        "[4] Schick, T., et al. (2023). Toolformer: Language models can teach themselves to "
        "use tools. Advances in Neural Information Processing Systems, 36.",
        "[5] Greshake, K., et al. (2023). Not what you've signed up for: Compromising "
        "real-world LLM-integrated applications with indirect prompt injection. Proceedings "
        "of the 16th ACM Workshop on Artificial Intelligence and Security, 79–90.",
        "[6] Ji, Z., et al. (2023). Survey of hallucination in natural language generation. "
        "ACM Computing Surveys, 55(12), Article 248.",
        "[7] Es, S., James, J., Espinosa-Anke, L., & Schockaert, S. (2024). RAGAS: Automated "
        "evaluation of retrieval augmented generation. Proceedings of EACL 2024: System "
        "Demonstrations, 150–158.",
        "[8] Malkov, Y. A., & Yashunin, D. A. (2020). Efficient and robust approximate nearest "
        "neighbor search using hierarchical navigable small world graphs. IEEE Transactions "
        "on Pattern Analysis and Machine Intelligence, 42(4), 824–836.",
        "[9] PostgreSQL Global Development Group. PostgreSQL 16 documentation: NOTIFY and "
        "SELECT locking clause. https://www.postgresql.org/docs/16/",
        "[10] Kane, A. pgvector: Open-source vector similarity search for Postgres. "
        "https://github.com/pgvector/pgvector",
    ], size=14, space_after=6)

    prs.save(str(OUT))
    print(f"wrote {OUT.relative_to(HERE.parent.parent)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    build(Path(sys.argv[1]))
